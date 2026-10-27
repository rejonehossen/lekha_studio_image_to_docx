"""
Table builder placeholder.

python-docx supports real tables; wiring them requires the structured_content
schema to include table blocks. Not currently used by docx_builder; kept as a
self-contained helper for future extension.
"""
from docx.shared import Pt


def add_table(doc, rows, style: str = "Light Grid Accent 1"):
    if not rows or not rows[0]:
        return None
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = style
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            cell.text = str(cell) if cell is not None else ""
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
    return table
